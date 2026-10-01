# find_scint_event.py

import numpy as np
import datetime as dt
from openpyxl import load_workbook

#wb = load_workbook('2017_polar_event_catalog.xlsx')
#ws = wb['Sheet1']
#date = [cell.value for cell in ws['A'][1:]]
#station = [cell.value for cell in ws['B'][1:]]
#prn = [int(cell.value) for cell in ws['C'][1:]]
#tow = [int(cell.value) for cell in ws['D'][1:]]
#sigphi = [float(cell.value) for cell in ws['G'][1:]]

wb = load_workbook('2017_aur_event_catalog.xlsx')
ws = wb['Sheet1']
date = [cell.value for cell in ws['B'][1:]]
station = [cell.value for cell in ws['C'][1:]]
prn = [int(cell.value) for cell in ws['D'][1:]]
tow = [int(cell.value) for cell in ws['E'][1:]]
sigphi = [float(cell.value) for cell in ws['K'][1:]]

#date = [np.datetime64(d[0:4]+'-'+d[4:6]+'-'+d[6:8]) for d in date]

# From Chintin
def parse_time(date, tow):
    hour_utc = int((tow % 86400) / 3600)
    minute = (((tow % 86400) / 3600) - int((tow % 86400) / 3600)) * 60
    second = int((minute - int(minute)) * 60)
    minute = int(minute)

    #time = dt.datetime.strptime(date, '%Y%m%d').replace(hour=hour_utc, minute=minute, second=second)
    time = dt.datetime.strptime(date, '%Y%m%d').replace(hour=hour_utc, minute=minute, second=0)
    return np.datetime64(time.isoformat())

time = [parse_time(d,t) for d, t in zip(date, tow)]


#date_ae = self.date
#date_ae = date_ae.replace(hour=hour_utc, minute=minute, second=second)
## print(kp_data)
#date_str = date_ae.strftime("%Y-%m-%d %H:%M:00")
#print(date_ae)
# 
#
#print(time)

scint_date = [d for d,s in zip(date, station) if s == 'san']
scint_time = [t for t,s in zip(time, station) if s == 'san']
scint_prn = [p for p,s in zip(prn, station) if s == 'san']
scint_sigphi = [sp for sp,s in zip(sigphi, station) if s == 'san']

#scint_date = [d for d,s in zip(date, station) if s == 'res']
#scint_time = [t for t,s in zip(time, station) if s == 'res']
#scint_prn = [p for p,s in zip(prn, station) if s == 'res']
#scint_sigphi = [sp for sp,s in zip(sigphi, station) if s == 'res']

#print(sigphi)

#conv = {
#    0: lambda x: x,  # conversion fn for column 0
#    1: lambda x: int(x[1:]),  # conversion fn for column 1
#    2: lambda x: np.datetime64(x),  # conversion fn for column 2
#}
#insitu, prn, time = np.loadtxt('insitu_gnss_conj_res_2017.txt', skiprows=1, usecols=[0,2,4], converters=conv, unpack=True)

conj_sat = list()
conj_prn = list()
conj_time = list()
conj_date = list()
conj_sep = list()
with open('insitu_gnss_conj_san_2017.txt', 'r') as f:
    next(f)
    for line in f:
        spltln = line.split()
        conj_sat.append(spltln[0]+spltln[1])
        conj_prn.append(int(spltln[3][1:]))
        conj_time.append(np.datetime64(spltln[5]+'T'+spltln[6]))
        conj_date.append(np.datetime64(spltln[5]))
        conj_sep.append(float(spltln[7]))

print(len(scint_date), len(conj_date))

for cd, ct, cp, cs, ce in zip(conj_date, conj_time, conj_prn, conj_sat, conj_sep):
    for sd, st, sp, ss in zip(scint_date, scint_time, scint_prn, scint_sigphi):
        if (ct == st) and (cp == sp):
            print(ct, cs, cp, sp, ss, ce)

#print(time)

#for sn in sheetnames:
#    print(sn)
#    ws = wb[sn]

    # Column A:	Year (2014-2019)
    # Column B:	day of year (1-366)
    # Column C:	frequency (1 for L1, 2 for L2C)
    # Column D:	GPS satellite PRN (1-32)
    # Column E:	Scintillation start UT hour (0-23)
    # Column F:	Scintillation start UT minute (0-59)
    # Column G:	Scintillation end UT hour (0-23)
    # Column H:	Scintillation end UT minute (0-59)
    # Column I:	number of receivers operational (4-6)

#    # pull colums describing start/end times skiping first row (heading)
#    year = [int(cell.value) for cell in ws['A'][1:]]
#    doy = [int(cell.value) for cell in ws['B'][1:]]
#    start_hour = [int(cell.value) for cell in ws['E'][1:]]
#    start_minute = [int(cell.value) for cell in ws['F'][1:]]
#    end_hour = [int(cell.value) for cell in ws['G'][1:]]
#    end_minute = [int(cell.value) for cell in ws['H'][1:]]
#
#    pfisr_experiment = ['PFISR Experiment']
#    pfisr_mode = ['PFISR Mode']
#    swarmA_st = ['Swarm A Start Time']
#    swarmA_et = ['Swarm A End Time']
#    swarmB_st = ['Swarm B Start Time']
#    swarmB_et = ['Swarm B End Time']
#    swarmC_st = ['Swarm C Start Time']
#    swarmC_et = ['Swarm C End Time']
#
#    for y, d, sh, sm, eh, em in zip(year, doy, start_hour, start_minute, end_hour, end_minute):
#
#        starttime = dt.time(sh,sm)
#        endtime = dt.time(eh,em)
#
#        startdate = dt.date(y,1,1)+dt.timedelta(days=d-1)
#        enddate = dt.date(y,1,1)+dt.timedelta(days=d-1)

